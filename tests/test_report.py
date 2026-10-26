import hashlib
import json
import zipfile

from openpyxl import load_workbook

from soc2kit.report import write_pack
from soc2kit.runner import scan


def _scan(session):
    session.client("s3").create_bucket(Bucket="example-bucket")
    return scan(session, regions=["us-east-1"], check_ids=["DATA-02", "DATA-05", "LOG-01", "NET-01", "LOG-08"])


def test_pack_contents_and_manifest(session, tmp_path):
    summary = write_pack(_scan(session), str(tmp_path))
    for name in ("report.json", "report.html", "workbook.xlsx", "manifest.json"):
        assert (tmp_path / name).exists(), name
    manifest = json.loads((tmp_path / "manifest.json").read_text())
    for rel, digest in manifest["files"].items():
        assert hashlib.sha256((tmp_path / rel).read_bytes()).hexdigest() == digest, rel
    assert "evidence/LOG-01.json" in manifest["files"]
    zips = list(tmp_path.glob("*.zip"))
    assert len(zips) == 1 and "manifest.json" in zipfile.ZipFile(zips[0]).namelist()
    assert summary["counts"]["FAIL"] >= 2


def test_workbook_sheets(session, tmp_path):
    write_pack(_scan(session), str(tmp_path), make_zip=False)
    wb = load_workbook(tmp_path / "workbook.xlsx")
    assert wb.sheetnames == ["Summary", "Automated Checks", "Failed Resources", "Manual Controls", "Criteria Coverage"]
    assert wb["Manual Controls"].max_row > 20


def test_untrusted_resource_names_are_neutralised(session, tmp_path):
    # Log group names are attacker-influenced strings that land in the HTML and XLSX.
    logs = session.client("logs")
    logs.create_log_group(logGroupName="=HYPERLINK(\"http://evil\")")
    logs.create_log_group(logGroupName="<script>alert(1)</script>")
    write_pack(scan(session, regions=["us-east-1"], check_ids=["LOG-08"]), str(tmp_path), make_zip=False)

    html = (tmp_path / "report.html").read_text()
    assert "<script>alert(1)</script>" not in html

    ws = load_workbook(tmp_path / "workbook.xlsx")["Failed Resources"]
    values = [c.value for row in ws.iter_rows(min_row=2) for c in row if isinstance(c.value, str)]
    assert not any(v.startswith("=") for v in values)
