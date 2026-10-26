"""SOC 2 workbook (XLSX): summary, automated results, findings, manual controls, criteria coverage."""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

_HEADER = PatternFill("solid", fgColor="1F2937")
_FILLS = {
    "PASS": PatternFill("solid", fgColor="D1FAE5"),
    "FAIL": PatternFill("solid", fgColor="FEE2E2"),
    "ERROR": PatternFill("solid", fgColor="FEF3C7"),
    "N/A": PatternFill("solid", fgColor="F3F4F6"),
}


def _safe(value):
    """Neutralise spreadsheet formula injection from resource names coming out of AWS."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list], widths: list[int],
           status_col: int | None = None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEADER
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        ws.append([_safe(v) for v in row])
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        if status_col is not None:
            st = row[status_col].value
            if st in _FILLS:
                row[status_col].fill = _FILLS[st]
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    return ws


def write_workbook(summary: dict, path: str) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    c = summary["counts"]
    _sheet(wb, "Summary", ["Field", "Value"], [
        ["AWS account", summary["account_id"]],
        ["Scanned as", summary["principal_arn"]],
        ["Regions", ", ".join(summary["regions"])],
        ["Scan started (UTC)", summary["started_at"]],
        ["Scan finished (UTC)", summary["finished_at"]],
        ["Automated score", f"{summary['score']}% of evaluated controls passing"],
        ["Controls: pass / fail / error / N/A", f"{c['PASS']} / {c['FAIL']} / {c['ERROR']} / {c['N/A']}"],
        ["Region-level results", str(sum(summary["result_counts"].values()))],
        ["Tool", f"{summary['tool']['name']} {summary['tool']['version']}"],
        ["Note", summary["disclaimer"]],
    ], [28, 100])

    _sheet(wb, "Controls",
           ["Check", "Title", "Criteria", "Severity", "Status", "Result", "Remediation", "Evidence files"],
           [[c["check_id"], c["title"], ", ".join(c["criteria"]), c["severity"], c["status"], c["summary"],
             c["remediation"], ", ".join(c["evidence_files"][:5]) + (" ..." if len(c["evidence_files"]) > 5 else "")]
            for c in summary["checks"]],
           [10, 42, 14, 11, 9, 46, 60, 40], status_col=4)

    _sheet(wb, "Region Detail",
           ["Check", "Title", "Region", "Status", "Result", "Evidence file", "Evaluated (UTC)"],
           [[r["check_id"], r["title"], r["region"], r["status"],
             r["summary"] + (f" ({r['error']})" if r["error"] else ""), r["evidence_file"], r["evaluated_at"]]
            for r in summary["results"]],
           [10, 42, 14, 9, 40, 34, 22], status_col=3)

    findings = [[r["check_id"], r["region"], f["resource"], f["status"], f["detail"], r["remediation"]]
                for r in summary["results"] for f in r["findings"] if f["status"] == "FAIL"]
    _sheet(wb, "Failed Resources", ["Check", "Region", "Resource", "Status", "Detail", "Remediation"],
           findings, [10, 14, 60, 9, 40, 60], status_col=3)

    ws = _sheet(wb, "Manual Controls",
                ["ID", "Criterion", "Control", "Evidence to collect", "Frequency", "Owner", "Status",
                 "Evidence link", "Notes"],
                [[m["id"], m["criterion"], m["control"], m["evidence"], m["frequency"], "", "Not started", "", ""]
                 for m in summary["manual_controls"]],
                [8, 10, 46, 60, 18, 18, 16, 30, 30])
    dv = DataValidation(type="list", formula1='"Not started,In progress,Evidence collected,Not applicable"',
                        allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"G2:G{len(summary['manual_controls']) + 1}")

    _sheet(wb, "Criteria Coverage",
           ["Criterion", "Description", "Automated checks", "Manual controls", "Pass", "Fail", "Error", "N/A"],
           [[r["criterion"], r["name"], ", ".join(r["automated_checks"]), ", ".join(r["manual_controls"]),
             r["pass"], r["fail"], r["error"], r["na"]] for r in summary["criteria"]],
           [10, 50, 30, 24, 7, 7, 7, 7])

    wb.save(path)
